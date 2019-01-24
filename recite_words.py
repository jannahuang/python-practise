import random
import docx
import openpyxl


def word_dict():
    # 从 word 导入内容，返回字典，单词是 key，中文是 value
    word_file = docx.Document('历年考研考生最易错的73组考研单词(打印版).docx')
    word_words = {}
    for para in word_file.paragraphs:
        if ' ' in para.text:
            lines = para.text.split(' ')[1:]
            index = []
            for i in lines:
                if len(i) >= 2 and i[1] in 'abcdefghijklmnopqrstuvwxyz':
                    index.append(lines.index(i))
            for i in range(len(index)):
                if i != len(index) - 1:
                    word_words[str(lines[index[i]])] = ''.join(lines[i + 1:index[i + 1]])
                else:
                    word_words[str(lines[index[i]])] = ''.join(lines[index[i] + 1:])
    return word_words


def excel_dict():
    # 从 excel 导入内容，返回字典，单词是 key，中文是 value
    excel_file = openpyxl.load_workbook('最常用2000个英语单词-EXCEL版.xlsx')
    sheet = excel_file['按字母排序']
    excel_words = {}
    for i in range(2, 2002):
        cell_of_word = 'B' + str(i)
        cell_of_mean = 'D' + str(i)
        excel_words[str(sheet[cell_of_word].value)] = sheet[cell_of_mean].value
    return excel_words


def load_dicts():
    # 合并 word 和 excel 的单词字典
    return dict(word_dict(), **excel_dict())


def random_word(dictionary):
    # 随机从字典抽出 1 个单词
    word = list(dictionary.keys())
    word = str(random.sample(word, 1))
    word = word.replace("[", "")
    word = word.replace("]", "")
    word = word.strip("'")
    print(dictionary[word])
    answer = str(input('word: '))
    if answer == word:
        print('right')
    else:
        print('wrong, the word is {}'.format(word))


def recite_words():
    # 设置背诵单词数
    number = int(input('今天要背诵的单词数量：'))
    for i in range(number):
        random_word(words)
    print('今天的任务完成，棒棒哒！')


if __name__ == '__main__':
    # 程序入口
    words = load_dicts()
    recite_words()
